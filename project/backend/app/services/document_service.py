import os
import uuid
from pathlib import Path

import numpy as np
import pandas as pd
import pdfplumber
import pymupdf
import pytesseract

from PIL import Image
from docx import Document
from pptx import Presentation
from langchain_text_splitters import RecursiveCharacterTextSplitter

from app.config import settings


BASE_DIR = Path("data")
UPLOAD_DIR = BASE_DIR / "uploads"
TEXT_DIR = BASE_DIR / "text"
IMAGE_DIR = BASE_DIR / "images"
TABLE_DIR = BASE_DIR / "tables"
OCR_DIR = BASE_DIR / "ocr"


for directory in [
    UPLOAD_DIR,
    TEXT_DIR,
    IMAGE_DIR,
    TABLE_DIR,
    OCR_DIR
]:
    directory.mkdir(
        parents=True,
        exist_ok=True
    )


SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".doc",
    ".pptx",
    ".ppt",
    ".xlsx",
    ".xls",
    ".csv",
    ".txt",
    ".jpg",
    ".jpeg",
    ".png",
    ".webp",
    ".bmp",
    ".tiff",
    ".tif"
}


text_splitter = RecursiveCharacterTextSplitter(
    chunk_size=settings.CHUNK_SIZE,
    chunk_overlap=settings.CHUNK_OVERLAP,
    length_function=len
)


def safe_name(value):
    return (
        str(value)
        .replace("/", "_")
        .replace("\\", "_")
        .replace(" ", "_")
    )


def table_to_text(rows):
    lines = []

    for row in rows:
        cells = [
            ""
            if cell is None
            else str(cell)
            .strip()
            .replace("\n", " ")
            for cell in row
        ]

        lines.append(" | ".join(cells))

    return "\n".join(lines)


def dataframe_to_text(dataframe):

    dataframe = dataframe.dropna(
        how="all"
    ).dropna(
        axis=1,
        how="all"
    )

    if dataframe.empty:
        return ""

    return dataframe.astype(str).to_csv(
        index=False,
        sep="|"
    )


def add_chunks(
    items,
    text,
    document,
    location,
    item_type="text",
    folder="text"
):

    if not str(text).strip():
        return

    chunks = text_splitter.split_text(
        str(text)
    )

    for chunk_number, chunk in enumerate(chunks):

        file_path = (
            BASE_DIR
            / folder
            / (
                f"{document['document_id']}_"
                f"{item_type}_"
                f"{safe_name(location)}_"
                f"{chunk_number}.txt"
            )
        )

        file_path.write_text(
            chunk,
            encoding="utf-8"
        )

        items.append({
            "document_id": document["document_id"],
            "file_name": document["file_name"],
            "page": location,
            "location": location,
            "type": item_type,
            "text": chunk,
            "path": str(file_path)
        })


def extract_pdf(document, items):

    pdf = pymupdf.open(
        document["filepath"]
    )

    with pdfplumber.open(
        document["filepath"]
    ) as plumber:

        for page_number in range(len(pdf)):

            page = pdf[page_number]

            location = (
                f"Page {page_number + 1}"
            )

            # Native text
            add_chunks(
                items,
                page.get_text(),
                document,
                location
            )

            # OCR
            ocr_path = (
                OCR_DIR
                / (
                    f"{document['document_id']}"
                    f"_page_{page_number}.png"
                )
            )

            page.get_pixmap(
                dpi=200
            ).save(ocr_path)

            try:
                ocr_text = pytesseract.image_to_string(
                    Image.open(ocr_path)
                )

                add_chunks(
                    items,
                    ocr_text,
                    document,
                    location,
                    item_type="ocr",
                    folder="ocr"
                )
            except Exception:
                pass

            # Embedded images
            for image_number, image in enumerate(
                page.get_images(full=True)
            ):

                xref = image[0]

                try:
                    pixmap = pymupdf.Pixmap(
                        pdf,
                        xref
                    )

                    if pixmap.n - pixmap.alpha > 3:
                        pixmap = pymupdf.Pixmap(
                            pymupdf.csRGB,
                            pixmap
                        )

                    image_path = (
                        IMAGE_DIR
                        / (
                            f"{document['document_id']}_"
                            f"page_{page_number}_"
                            f"image_{image_number}.png"
                        )
                    )

                    pixmap.save(
                        image_path
                    )

                    items.append({
                        "document_id":
                            document["document_id"],
                        "file_name":
                            document["file_name"],
                        "page":
                            location,
                        "location":
                            location,
                        "type":
                            "image",
                        "path":
                            str(image_path)
                    })

                except Exception:
                    pass

            # Tables
            pdf_page = plumber.pages[
                page_number
            ]

            tables = (
                pdf_page.extract_tables()
                or []
            )

            for table_number, table in enumerate(
                tables
            ):

                table_location = (
                    f"{location}, "
                    f"Table {table_number + 1}"
                )

                add_chunks(
                    items,
                    table_to_text(table),
                    document,
                    table_location,
                    item_type="table",
                    folder="tables"
                )

    pdf.close()


def extract_docx(document, items):

    word_doc = Document(
        document["filepath"]
    )

    paragraph_text = "\n".join(
        paragraph.text
        for paragraph in word_doc.paragraphs
    )

    add_chunks(
        items,
        paragraph_text,
        document,
        "Document text"
    )

    # Tables
    for table_number, table in enumerate(
        word_doc.tables,
        start=1
    ):

        rows = [
            [
                cell.text
                for cell in row.cells
            ]
            for row in table.rows
        ]

        add_chunks(
            items,
            table_to_text(rows),
            document,
            f"Table {table_number}",
            item_type="table",
            folder="tables"
        )

    # Embedded images
    for relationship in word_doc.part.rels.values():

        if "image" not in relationship.target_ref:
            continue

        try:
            image_data = (
                relationship.target_part.blob
            )

            image_path = (
                IMAGE_DIR
                / (
                    f"{document['document_id']}_"
                    f"word_image_"
                    f"{uuid.uuid4().hex}.png"
                )
            )

            image_path.write_bytes(
                image_data
            )

            items.append({
                "document_id":
                    document["document_id"],
                "file_name":
                    document["file_name"],
                "page":
                    "Word document",
                "location":
                    "Word document",
                "type":
                    "image",
                "path":
                    str(image_path)
            })

        except Exception:
            pass


def extract_pptx(document, items):

    presentation = Presentation(
        document["filepath"]
    )

    for slide_number, slide in enumerate(
        presentation.slides,
        start=1
    ):

        slide_text = []
        slide_tables = []

        for shape in slide.shapes:

            if getattr(
                shape,
                "text",
                ""
            ).strip():

                slide_text.append(
                    shape.text
                )

            if getattr(
                shape,
                "has_table",
                False
            ):

                rows = [
                    [
                        cell.text
                        for cell in row.cells
                    ]
                    for row in shape.table.rows
                ]

                slide_tables.append(rows)

            # Images
            if getattr(
                shape,
                "shape_type",
                None
            ) == 13:

                try:
                    image = shape.image

                    image_path = (
                        IMAGE_DIR
                        / (
                            f"{document['document_id']}_"
                            f"slide_{slide_number}_"
                            f"image_{uuid.uuid4().hex}."
                            f"{image.ext}"
                        )
                    )

                    image_path.write_bytes(
                        image.blob
                    )

                    items.append({
                        "document_id":
                            document["document_id"],
                        "file_name":
                            document["file_name"],
                        "page":
                            f"Slide {slide_number}",
                        "location":
                            f"Slide {slide_number}",
                        "type":
                            "image",
                        "path":
                            str(image_path)
                    })

                except Exception:
                    pass

        add_chunks(
            items,
            "\n".join(slide_text),
            document,
            f"Slide {slide_number}"
        )

        for table_number, rows in enumerate(
            slide_tables,
            start=1
        ):

            location = (
                f"Slide {slide_number}, "
                f"Table {table_number}"
            )

            add_chunks(
                items,
                table_to_text(rows),
                document,
                location,
                item_type="table",
                folder="tables"
            )


def extract_excel(document, items):

    sheets = pd.read_excel(
        document["filepath"],
        sheet_name=None
    )

    for sheet_name, dataframe in sheets.items():

        add_chunks(
            items,
            dataframe_to_text(dataframe),
            document,
            f"Sheet {sheet_name}",
            item_type="table",
            folder="tables"
        )


def extract_csv(document, items):

    dataframe = pd.read_csv(
        document["filepath"]
    )

    add_chunks(
        items,
        dataframe_to_text(dataframe),
        document,
        "CSV table",
        item_type="table",
        folder="tables"
    )


def extract_txt(document, items):

    text = Path(
        document["filepath"]
    ).read_text(
        encoding="utf-8",
        errors="ignore"
    )

    add_chunks(
        items,
        text,
        document,
        "Text file"
    )


def extract_image(document, items):

    image_path = document["filepath"]

    # Original image
    items.append({
        "document_id":
            document["document_id"],
        "file_name":
            document["file_name"],
        "page":
            "Image",
        "location":
            "Image",
        "type":
            "image",
        "path":
            image_path
    })

    # OCR
    try:

        ocr_text = pytesseract.image_to_string(
            Image.open(image_path)
        )

        add_chunks(
            items,
            ocr_text,
            document,
            "Image OCR",
            item_type="ocr",
            folder="ocr"
        )

    except Exception:
        pass


EXTRACTORS = {
    ".pdf": extract_pdf,
    ".docx": extract_docx,
    ".pptx": extract_pptx,
    ".xlsx": extract_excel,
    ".xls": extract_excel,
    ".csv": extract_csv,
    ".txt": extract_txt,

    ".jpg": extract_image,
    ".jpeg": extract_image,
    ".png": extract_image,
    ".webp": extract_image,
    ".bmp": extract_image,
    ".tiff": extract_image,
    ".tif": extract_image,
}


def process_document(document):

    extension = document["extension"]

    if extension not in EXTRACTORS:
        raise ValueError(
            f"Unsupported file type: {extension}"
        )

    items = []

    EXTRACTORS[
        extension
    ](
        document,
        items
    )

    return items